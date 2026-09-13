"""3-5 framing slides for the HavenIQ meeting: intro, context, architecture, the call we will demo, what it costs to run."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
INK=RGBColor(0x14,0x1c,0x2b); MUTED=RGBColor(0x62,0x6c,0x7a); LINE=RGBColor(0xd9,0xde,0xe6); ACC=RGBColor(0x0f,0x7c,0x6b); SOFT=RGBColor(0xef,0xf7,0xf5); WHITE=RGBColor(255,255,255); WARM=RGBColor(0xfd,0xf3,0xe7)
HD="Avenir Next"; TX="Avenir Next"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
def text(s,x,y,w,h,t,size,bold=False,color=INK,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,font=TX,italic=False):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    lines=t if isinstance(t,list) else [t]
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        if i: p.space_before=Pt(6)
        r=p.add_run(); r.text=ln; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.name=font; r.font.color.rgb=color
    return tb
def box(s,x,y,w,h,title,sub="",fill=WHITE,line=LINE,tsize=13,ssize=10,accent=None):
    b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); b.adjustments[0]=0.08
    b.fill.solid(); b.fill.fore_color.rgb=fill; b.line.color.rgb=line; b.line.width=Pt(1); b.shadow.inherit=False
    tf=b.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.margin_left=tf.margin_right=Inches(0.12); tf.margin_top=tf.margin_bottom=Inches(0.06)
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=title; r.font.size=Pt(tsize); r.font.bold=True; r.font.name=HD; r.font.color.rgb=accent or INK
    if sub:
        p2=tf.add_paragraph(); p2.alignment=PP_ALIGN.CENTER; r2=p2.add_run(); r2.text=sub; r2.font.size=Pt(ssize); r2.font.name=TX; r2.font.color.rgb=MUTED
    return b
def arrow(s,x1,y1,x2,y2,label="",color=MUTED):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); c.line.color.rgb=color; c.line.width=Pt(1.5)
    c.line._get_or_add_ln().append(__import__("pptx.oxml").oxml.parse_xml('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" type="triangle"/>'))
    if label: text(s,(x1+x2)/2-0.9,(y1+y2)/2-0.28,1.8,0.3,label,9,False,MUTED,PP_ALIGN.CENTER)
def chrome(s,n,total=5):
    text(s,0.5,7.0,8,0.3,"HavenIQ voice support  ·  LiveKit  ·  Dean Peterson",10,False,MUTED); text(s,12.0,7.0,0.83,0.3,f"{n}/{total}",10,False,MUTED,PP_ALIGN.RIGHT)
# 1 intro
s=prs.slides.add_slide(blank); chrome(s,1)
text(s,0.7,1.6,11.9,1.2,"A voice support agent for HavenIQ",44,True,INK,font=HD)
text(s,0.7,2.9,10.5,1.2,"Dean Peterson, Solutions Architect, LiveKit. Thirty minutes: the situation, the design, a live call, your questions throughout.",20,False,MUTED)
text(s,0.7,4.5,11.9,0.5,"Before we start, a quick round: who you are and what you want out of a support agent.",16,False,INK)
for i,(who,role) in enumerate([("Mike","Director of AI Platform"),("Ahmad","Director of Customer Support"),("Lucy","VP of Finance")]):
    box(s,0.7+i*3.9,5.15,3.6,0.95,who,role,SOFT,SOFT,15,11,ACC)
# 2 context
s=prs.slides.add_slide(blank); chrome(s,2)
text(s,0.7,0.7,11.9,0.8,"Where HavenIQ is today",36,True,INK,font=HD)
text(s,0.7,1.55,11.9,0.6,"Connected thermostats, cameras, locks and sensors, plus a cloud video subscription. Every customer has an account, registered devices, a subscription and orders.",15,False,MUTED)
cols=[("What callers ask","Where is my order. Is my thermostat online. What is the temperature in the living room. Why did my subscription lapse. Most of it is a lookup.",SOFT),
      ("What it costs today","Every call starts with identity and context gathering. Routine questions take a person's time before the real problem surfaces. Handoffs lose what was already said.",WARM),
      ("What we are building","A voice agent that greets by name, answers the routine questions from HavenIQ's own systems, and hands a person a summary and the recording when it should not decide alone.",SOFT)]
for i,(t,b,f) in enumerate(cols):
    box(s,0.7+i*4.05,2.5,3.85,3.6,"","",f,f,16,11,ACC); text(s,0.95+i*4.05,2.75,3.35,0.5,t,16,True,ACC,font=HD); text(s,0.95+i*4.05,3.3,3.35,2.6,b,13,False,INK)
text(s,0.7,6.35,11.9,0.5,"Success looks like: shorter calls, fewer repeated questions, and a human who joins already knowing the story.",14,True,INK)
# 3 architecture
s=prs.slides.add_slide(blank); chrome(s,3)
text(s,0.7,0.55,11.9,0.7,"Architecture",30,True,INK,font=HD)
text(s,0.7,1.2,11.9,0.4,"LiveKit carries the voice; HavenIQ's platform runs the agent. ",13,False,MUTED)
text(s,0.7,1.5,11.9,0.4,"Media plane in LiveKit Cloud; agent, tools, tickets and recordings on HavenIQ's OpenShift cluster, all declared in git.",13,False,MUTED)
box(s,0.6,2.5,2.0,1.1,"Caller","browser, or phone via SIP",WHITE,LINE,13,10)
box(s,3.2,2.5,2.2,1.1,"LiveKit Cloud","rooms, WebRTC, egress, inference (STT, TTS)",SOFT,ACC,13,10,ACC)
box(s,6.1,2.2,2.6,1.7,"Voice worker","LiveKit Agents SDK on OpenShift: listen, think, speak, hand off",WHITE,LINE,13,10)
box(s,9.4,2.5,3.3,1.1,"Model desk","LLM behind a guardrail gateway; model is a config change",WHITE,LINE,13,10)
box(s,6.1,4.55,2.6,1.15,"MCP gateway","every tool call authorized and logged; no credentials in the agent",WARM,LINE,13,10)
box(s,3.0,5.95,2.6,1.0,"HavenIQ cloud","accounts, devices, orders (mock)",WHITE,LINE,12,10)
box(s,6.1,5.95,2.6,1.0,"CRM bridge","tickets, notes, handoff page",WHITE,LINE,12,10)
box(s,9.2,5.95,2.0,1.0,"Zammad","the support team's desk",WHITE,LINE,12,10)
box(s,9.4,4.55,3.3,1.15,"Object store (NooBaa)","call recording, attached to the ticket",WHITE,LINE,12,10)
box(s,0.6,4.55,2.0,1.15,"Human agent","joins the same room from the desk with the summary",SOFT,LINE,12,10)
arrow(s,2.6,3.05,3.2,3.05,"audio"); arrow(s,5.4,3.05,6.1,3.05,"dispatch"); arrow(s,8.7,3.05,9.4,3.05,"turns")
arrow(s,7.4,3.9,7.4,4.55,"tool calls"); arrow(s,6.1,5.12,4.3,5.95,""); arrow(s,7.4,5.7,7.4,5.95,""); arrow(s,8.7,6.45,9.2,6.45,"API")
arrow(s,5.4,2.9,9.4,4.8,"egress → recording"); arrow(s,2.6,5.12,3.2,3.3,"joins room")
# 4 the call
s=prs.slides.add_slide(blank); chrome(s,4)
text(s,0.7,0.7,11.9,0.8,"The call you are about to hear",36,True,INK,font=HD)
steps=[("1  Greet by name","The caller's number arrives with the call. The agent looks the account up through the gateway before the first word and greets Dana by name."),
       ("2  Answer from the source","Order status, thermostat status, living-room temperature: each answer is a governed tool call to HavenIQ's cloud, never a guess."),
       ("3  Hand off with the story","When Dana asks for a person, the agent opens the ticket with a summary, pages the desk, and a human joins the same room already briefed."),
       ("4  Close the loop","The recording lands in HavenIQ's own object store and is attached to the ticket, so the follow-up has both the summary and the audio.")]
for i,(t,b) in enumerate(steps):
    y=1.75+i*1.25; box(s,0.7,y,3.2,1.0,t,"",SOFT,SOFT,15,10,ACC); text(s,4.1,y+0.12,8.5,0.9,b,14,False,INK)
text(s,0.7,6.85,11.9,0.4,"Interrupt any time; the agent handles interruptions the same way you just did.",12,False,MUTED,italic=True)
# 5 running it
s=prs.slides.add_slide(blank); chrome(s,5)
text(s,0.7,0.7,11.9,0.8,"What it takes to run",36,True,INK,font=HD)
cols=[("Latency budget","Speech to text, one model turn with tools, text to speech. The model turn is the variable; tool calls are local to the cluster. Target under two seconds to first word.",SOFT),
      ("Cost per minute","LiveKit media and inference are metered per minute; the model per token; the rest is cluster capacity you already own. Routine calls cost cents; a human's time is the expensive part we protect.",WARM),
      ("Control","Every tool the agent can use is registered and authorized per call, with an audit line. The agent holds no credentials. Swapping the model is a configuration change, not a rebuild.",SOFT),
      ("What is mock today","Accounts, devices and orders are a seeded store. Replace it with your real APIs behind the same gateway; the agent's calls do not change.",WHITE)]
for i,(t,b,f) in enumerate(cols):
    box(s,0.7+i*3.05,1.8,2.9,4.6,"","",f,LINE if f==WHITE else f,15,10,ACC); text(s,0.9+i*3.05,2.05,2.5,0.5,t,15,True,ACC,font=HD); text(s,0.9+i*3.05,2.6,2.5,3.7,b,12.5,False,INK)
text(s,0.7,6.6,11.9,0.5,"Questions welcome now, or after the call.",14,True,INK)
prs.save("haveniq-meeting.pptx"); print("saved haveniq-meeting.pptx", len(prs.slides), "slides")
